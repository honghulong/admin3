from copy import deepcopy
from docx import Document
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = r'C:\Users\Administrator\Desktop\ai套餐选择.docx'
doc = Document(path)
before = [etree.tostring(e) for e in doc.element.body if e.tag != qn('w:sectPr')]

def para(text, style='Normal (Web)'):
    return doc.add_paragraph(text, style)

def heading(text, level=2):
    para(text, 'Heading ' + str(level))

def table(headers, rows, reference):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = reference.style
    t._tbl.replace(t._tbl.tblPr, deepcopy(reference._tbl.tblPr))
    for c, value in zip(t.rows[0].cells, headers):
        c.text = value
    for values in rows:
        for c, value in zip(t.add_row().cells, values):
            c.text = value
    for r in t.rows:
        for i, c in enumerate(r.cells):
            c.width = reference.rows[0].cells[i].width
            for p in c.paragraphs:
                p.style = reference.rows[0].cells[i].paragraphs[0].style
    for c in t.rows[0].cells:
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'D9EAF7')
        c._tc.get_or_add_tcPr().append(shade)
        for run in c.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    return t

heading('六、阿里云 Token Plan 企业套餐补充')
heading('推荐结论')
para('推荐 Token Plan 团队版标准坐席，作为约 150 元/人/月的企业采购选项。官方产品名称为“团队版”，适合需要统一管理成员、模型用量和登录权限的企业。价格与 TRAE 团队版处在同一预算区间。')
para('推荐理由概括：')
para('• 标准坐席限时 150 元/人/月，原价 198 元；每人每月含 25,000 Credits。150 元属于优惠价，购买与续费以订单页为准。')
para('• 提供团队后台和独立 API Key；支持 SSO、钉钉接入及成员用量分析。')
para('• 套餐提供模型调用额度，需要接入兼容的 AI 编程或智能体工具使用，工具自身收费需另行核实。')
heading('一、套餐对比')
para('数据来源：阿里云官方团队版概述与团队管理文档；核验日期：2026 年 10 月 10 日。')
table(['对比项', '标准坐席（推荐）', '高级坐席'], [
    ['单人月费', '限时 150 元；原价 198 元', '限时 550 元；原价 698 元'],
    ['每人月额度', '25,000 Credits', '100,000 Credits'],
    ['8 人月费', '优惠价 1,200 元；原价 1,584 元', '优惠价 4,400 元；原价 5,584 元'],
    ['8 人年度预算', '优惠价按月续费 14,400 元；原价 19,008 元', '优惠价按月续费 52,800 元；原价 67,008 元'],
    ['使用强度', '轻度 AI 辅助', '日常高频使用'],
    ['企业管理', '团队后台、席位管理、成员用量分析', '同标准坐席'],
    ['企业登录', 'SAML 2.0 SSO、钉钉接入', '同标准坐席'],
    ['数据使用', '承诺不使用对话数据训练模型', '同标准坐席'],
    ['接入方式', '成员独立 API Key + 套餐 Base URL', '同标准坐席'],
], doc.tables[2])
heading('二、方案分析')
heading('1. 成本效益分析', 3)
para('按优惠价，标准坐席单人每年预算为 1,800 元；8 人合计 1,200 元/月，比原文 TRAE 团队版 1,192 元/月高 8 元。两者费用接近，阿里云方案的选择重点是多模型接入和企业管理。年度数值为月费乘以 12 的预算测算，不代表已确认全年优惠或包年报价。')
heading('2. Credits 额度够不够用？', 3)
para('Credits 不是固定 Token 数或请求次数；消耗取决于模型、输入输出长度、思考模式和工具调用。建议先用标准坐席跑实际开发任务，再根据成员用量判断是否升级。额度耗尽后，没有共享用量包则会暂停服务。')
heading('3. 企业功能是否满足需求？', 3)
para('成员通过专属 API Key 使用模型，管理者可查看消耗并管理席位与登录方式，适合集中采购。该方案仅限兼容编程和智能体工具中的交互式使用，不可用于应用后端或自动化脚本；也不包含 TRAE、Codex 等产品的订阅权益。')
heading('三、相关链接')
table(['用途', '链接'], [
    ['官方套餐价格与额度', 'https://help.aliyun.com/zh/model-studio/token-plan-team-overview'],
    ['企业管理与 SSO 配置', 'https://help.aliyun.com/zh/model-studio/token-plan-team-management'],
    ['订阅与接入入口', 'https://bailian.console.aliyun.com/'],
], doc.tables[3])
heading('四、行动建议')
para('第一步：在购买页选择团队版标准坐席，核对单席成交价、购买数量及续费价格；先按月采购，按每人 150 元设置预算。')
para('第二步：为成员分配独立 API Key，接入现有兼容编程工具；需要统一登录时配置 SSO 或钉钉。')
para('第三步：观察一个月的真实用量和开发效果。额度充足者保持标准坐席；频繁耗尽者单独评估高级坐席，避免全员升级。')

current = [e for e in doc.element.body if e.tag != qn('w:sectPr')]
assert all(a == etree.tostring(b) for a, b in zip(before, current))
assert len(doc.tables) == 6
doc.save(path)
print('Appended Aliyun section; original body elements preserved; 6 tables verified.')
