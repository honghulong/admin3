from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path("docs/企业费用报销业务管理系统产品功能规划说明书.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT = "E8EEF5"
GRAY = "F2F4F7"
MUTED = "666666"
INK = "222222"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

def set_font(run, size=11, bold=False, color=INK, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for s, size, color, before, after in [
    ("Title", 24, INK, 0, 6),
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 11.5, DARK, 8, 4),
]:
    st = styles[s]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ["List Bullet", "List Number"]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.5)
    st.paragraph_format.first_line_indent = Inches(-0.25)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.15

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths))); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(w)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths[i])); tcW.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[i] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        shade(t.rows[0].cells[i], LIGHT)
        p = t.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(h), 9.5, True, DARK)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(val)), 9.2)
    set_table_geometry(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def p(text="", bold_prefix=None):
    para = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_font(para.add_run(bold_prefix), bold=True)
        set_font(para.add_run(text[len(bold_prefix):]))
    else:
        set_font(para.add_run(text))
    return para

def bullets(items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        set_font(para.add_run(item))

def nums(items):
    for item in items:
        para = doc.add_paragraph(style="List Number")
        set_font(para.add_run(item))

def callout(title, text):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    set_table_geometry(t, [9360])
    shade(t.cell(0, 0), GRAY)
    para = t.cell(0, 0).paragraphs[0]
    set_font(para.add_run(title + "："), 10.5, True, DARK)
    set_font(para.add_run(text), 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

# running header/footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(hp.add_run("企业费用报销业务管理系统 | 产品规划"), 8.5, color=MUTED)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("内部评审稿 · 2026年7月"), 8.5, color=MUTED)

# cover / masthead
doc.add_paragraph()
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(title.add_run("企业费用报销业务管理系统"), 24, True, INK)
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(18)
set_font(sub.add_run("产品功能规划说明书（职工提交、业务审批、财务审核、付款与归档）"), 13, color=MUTED)
table(["文档属性", "内容"], [
    ("文档类型", "产品功能规划 / PRD"),
    ("适用对象", "公司职工、部门负责人、项目负责人、财务人员、出纳、系统管理员、审计人员"),
    ("适用范围", "中国大陆一般企业日常费用报销；具体制度、额度和税务口径由公司配置并经财务确认"),
    ("版本", "V1.0（内部评审稿）"),
    ("编制日期", "2026年7月9日"),
], [1900, 7460])
callout("核心目标", "让职工少填、审批人看得懂、财务审得准、出纳付得清、审计查得到，并通过规则、发票唯一性和完整日志降低重复报销与违规风险。")

doc.add_heading("1. 文档定位与规划原则", level=1)
p("本说明书用于统一报销系统的产品范围、角色职责、业务流程、功能模块、数据要求和分期建设计划，可作为需求评审、开发排期、测试验收和制度落地的共同依据。它不是公司财务制度，也不替代会计、税务或档案专业判断。")
bullets([
    "业务闭环：覆盖事前申请（可选）、报销提交、业务审批、财务审核、付款、入账、归档和审计查询。",
    "一票一档：原始文件名用于展示，唯一存储名和附件ID用于存储、查重、预览、删除与追溯。",
    "规则前置：在提交前完成必填、额度、预算、发票重复和基础合规检查，减少审批往返。",
    "权责分离：申请、业务审批、财务审核、付款和系统配置按角色授权，关键操作留痕。",
    "配置优先：审批链、费用标准、预算维度、单据编号和会计映射均由后台配置，避免写死。",
])

doc.add_heading("2. 用户角色与职责", level=1)
table(["角色", "主要职责", "关键权限边界"], [
    ("职工/申请人", "创建报销单、上传凭证、确认OCR结果、提交、撤回、补正、查询进度", "仅操作本人或被授权代办的单据；已付款单据不可修改"),
    ("部门/项目审批人", "核实业务真实性、必要性、归属部门/项目和费用合理性", "不可修改原始凭证；审批动作和意见必须留痕"),
    ("财务审核人", "审核票据、制度、预算、税务处理、科目和成本归属", "可退回补正、调整财务字段；不得代替申请人修改业务事实"),
    ("出纳/付款人", "生成付款批次、复核收款信息、登记支付结果和回单", "原则上与财务审核分离；付款失败须回退处理"),
    ("财务主管", "处理例外、复核高风险/大额单据、关账与报表", "拥有规则豁免权时必须填写原因并形成审计记录"),
    ("系统管理员", "组织、角色、流程、字典、规则、接口和安全配置", "默认无财务审批权；不可篡改业务与审计记录"),
    ("审计/管理层", "只读查询、抽样、异常分析和导出", "按组织及数据范围授权，不参与日常流转"),
], [1200, 4400, 3760])

doc.add_heading("3. 端到端业务流程", level=1)
nums([
    "职工新建报销单，选择费用类型、部门/项目、事由、发生日期和收款方式。",
    "上传发票、行程单、支付截图等凭证；系统生成唯一附件ID和唯一存储名，保留原文件名。",
    "OCR提取发票号码、代码、日期、购销方、金额税额等信息，申请人确认或修正。",
    "系统执行必填、金额勾稽、发票重复、预算、费用标准、敏感词和附件完整性检查。",
    "申请人提交；系统根据组织、费用类型、金额、项目和例外条件自动匹配审批链。",
    "业务审批人审核真实性与合理性，可同意、退回补正、拒绝或转交（按制度配置）。",
    "财务审核人完成票据、税务、预算、科目和支付信息复核；必要时触发加签或财务主管复核。",
    "审批通过后进入待付款；出纳生成付款批次并登记银行流水、付款日期和失败原因。",
    "付款成功后生成或对接记账凭证，关联报销单、附件、审批记录和银行回单。",
    "系统按归档规则固化电子凭证及元数据，支持查询、下载、移交、备份和审计。"
])

doc.add_heading("4. 单据状态与流转规则", level=1)
table(["状态", "进入条件", "允许操作", "下一状态"], [
    ("草稿", "新建或自动暂存", "编辑、上传/删除附件、保存、作废、提交", "审批中/已作废"),
    ("审批中", "提交且校验通过", "审批、退回、拒绝；申请人按规则撤回", "财务审核/待补正/已拒绝"),
    ("待补正", "审批人或财务退回", "申请人修改指定内容并再次提交", "审批中"),
    ("财务审核", "业务审批完成", "财务通过、退回、拒绝、加签、标记风险", "待付款/待补正/已拒绝"),
    ("待付款", "全部审核通过", "加入付款批次、复核收款信息", "付款中/付款失败"),
    ("付款中", "已提交银行或线下支付", "查询结果、登记回单", "已付款/付款失败"),
    ("已付款", "支付成功", "查看、入账、冲销/退款申请", "已入账/冲销中"),
    ("已入账/已归档", "凭证生成且归档完成", "只读查询、审计导出、红冲关联", "终态或更正流程"),
    ("已拒绝/已作废", "审批拒绝或草稿作废", "只读、复制新建", "终态"),
], [1300, 2300, 3300, 2460])
callout("流转约束", "禁止直接删除已提交单据；更正、冲销、退款必须创建关联单据。任何退回、拒绝、撤回、豁免和财务字段调整均应记录操作者、时间、前后值及原因。")

doc.add_heading("5. 功能架构", level=1)
table(["功能域", "一期必备", "增强能力"], [
    ("报销工作台", "我的草稿、待办、已办、进度、消息提醒", "移动端、代办、批量处理、智能催办"),
    ("单据与附件", "分类表单、自动保存、唯一附件、预览下载", "电子凭证原文件接收、邮件/卡包导入"),
    ("OCR与发票", "字段识别、人工确认、重复提示", "验真、状态查询、税额拆分、连号/异常分析"),
    ("审批引擎", "按组织/金额/类型配置顺序审批", "条件分支、并行会签、加签、委托、超时升级"),
    ("财务审核", "合规清单、预算提示、科目/成本归属", "规则引擎、风险评分、自动凭证建议"),
    ("付款管理", "待付款清单、批次、支付结果、回单", "银企直联、失败重试、付款对账"),
    ("入账与归档", "凭证号登记、附件包、全链路查询", "ERP接口、电子凭证标准化、档案系统接口"),
    ("分析与配置", "费用报表、组织角色、字典、审批配置", "预算看板、异常洞察、规则模拟、开放API"),
], [1500, 4000, 3860])

doc.add_heading("6. 职工端功能说明", level=1)
doc.add_heading("6.1 新建与填写", level=2)
bullets([
    "支持日常费用、差旅、交通、招待、采购垫付等费用类型；不同类型显示不同字段和附件要求。",
    "自动带出申请人、部门、默认成本中心和收款账户；敏感信息脱敏展示。",
    "按费用明细录入日期、类别、金额、税额、项目、事由、参与人/客户等业务要素。",
    "支持保存草稿、自动保存、复制历史单据和从事前申请/借款单关联生成。",
])
doc.add_heading("6.2 凭证与OCR", level=2)
bullets([
    "每次上传生成附件ID和UUID存储名，保留原文件名、大小、类型、哈希值、上传人和时间。",
    "同名文件可重复上传，但系统以附件ID/哈希/发票关键字段区分；删除只作用于选中的唯一附件。",
    "OCR结果必须可查看、可确认；低置信度字段高亮，修改保留识别值与人工修正值。",
    "支持图片和PDF预览；原始电子凭证不得因转图片或压缩而丢失，版式文件与结构化数据宜同时保存。",
])
doc.add_heading("6.3 提交前检查与进度", level=2)
bullets([
    "阻断项：必填缺失、金额不平、无必要附件、确定重复发票、无有效审批路径。",
    "提示项：预算不足、超标准、发票抬头异常、开票日期超期、OCR低置信度，可按权限继续提交并说明。",
    "流程视图显示当前节点、已处理人、处理时间、退回原因和预计下一节点；支持站内消息及可选企业IM提醒。",
])

doc.add_heading("7. 审批人与财务端功能说明", level=1)
doc.add_heading("7.1 业务审批", level=2)
bullets([
    "待办按紧急程度、提交时间、金额和风险标签排序；详情同时展示业务摘要、费用明细、附件和历史意见。",
    "支持同意、退回补正、拒绝、加签/转交（配置项）；退回必须指定问题项和意见。",
    "审批时展示部门/项目预算占用、同类历史费用和超标说明，但不得替代审批人判断。",
])
doc.add_heading("7.2 财务审核", level=2)
table(["审核维度", "系统辅助", "财务处理"], [
    ("票据", "OCR字段、文件哈希、发票关键字段、重复状态", "确认真实性、有效性、抬头、日期和凭证完整性"),
    ("金额", "含税金额、税额、报销金额勾稽", "确认可报金额、不可报金额和税务处理"),
    ("制度", "费用标准、超标项、缺失附件、敏感场景", "通过、退回、拒绝或有理由豁免"),
    ("预算", "预算余额、占用与释放记录", "确认预算维度及超预算处理方式"),
    ("核算", "科目、成本中心、项目、税码建议", "确认或调整财务字段并保留调整日志"),
    ("付款", "收款账户校验、重复付款提示", "确认付款对象、金额和支付方式"),
], [1200, 3860, 4300])
doc.add_heading("7.3 付款、入账与归档", level=2)
bullets([
    "按公司、币种、账户和付款日期组建付款批次；批次具有制单、复核、支付、结果回写状态。",
    "记录银行流水号、付款日期、付款回单和失败原因；付款失败返回待处理而非静默结束。",
    "生成会计凭证或对接ERP，建立“报销单—发票/附件—审批记录—付款回单—记账凭证”关联。",
    "归档包保存原始电子凭证及元数据、业务单据、审批记录、记账凭证和必要的检索关系。",
])

doc.add_heading("8. 管理后台与配置", level=1)
bullets([
    "组织与权限：员工、部门、岗位、角色、数据范围、审批代理、离职交接。",
    "费用类型：字段模板、附件要求、费用标准、可用科目、预算维度和启停时间。",
    "审批流程：按公司/部门/项目/金额/费用类型配置节点、条件、会签、加签、超时和版本。",
    "合规规则：重复发票、抬头、日期、金额、预算、招待人数、差旅标准等；支持阻断/警告/记录三级。",
    "基础数据：供应商、项目、成本中心、会计科目、税码、币种、银行账户和单据编号规则。",
    "接口与运维：SSO、组织同步、企业IM、发票服务、预算/ERP、银企直联、档案系统、日志与告警。",
])

doc.add_heading("9. 核心数据、权限与审计要求", level=1)
table(["对象", "关键数据", "控制要求"], [
    ("报销单", "编号、申请人、组织、类型、事由、金额、状态、流程版本", "状态机控制；提交后业务事实变更留痕"),
    ("费用明细", "日期、类别、金额税额、项目、成本中心、说明", "合计勾稽；支持一单多明细"),
    ("附件/发票", "附件ID、原名、唯一名、哈希、URL、OCR、发票关键字段", "唯一定位、防重复、原件保存、访问鉴权"),
    ("审批记录", "节点、处理人、动作、意见、时间、委托/加签关系", "追加写入，不允许普通用户修改或删除"),
    ("预算记录", "预算维度、占用、释放、实际发生", "幂等处理，退回/作废时正确释放"),
    ("付款记录", "批次、账户、流水、状态、回单", "制单复核分离；防重复支付"),
    ("会计/档案", "凭证号、科目、归档包、移交状态", "可检索、可导出、可校验、可备份"),
], [1300, 4300, 3760])
bullets([
    "权限采用角色权限 + 数据范围 + 单据关系三层控制，附件下载接口必须执行同等鉴权。",
    "审计日志至少记录登录、查看敏感信息、导出、上传/删除附件、提交、审批、财务调整、付款和配置变更。",
    "敏感字段（身份证、银行卡等）加密存储、脱敏展示；导出和批量下载需授权并留痕。",
    "关键接口具备幂等键、重试、超时和对账机制，防止重复提交、重复付款或状态丢失。",
])

doc.add_heading("10. 合规与档案设计依据", level=1)
p("产品设计应支持经办、审核、审批等必要审签程序，防止电子会计凭证重复入账，并能够接收、读取、输出电子凭证及其元数据。电子档案还需关注真实、完整、可用、安全、备份、长期保管及检索关系。上述能力需要结合公司会计制度、税务处理和档案管理办法，由财务与档案责任部门最终确认。")
bullets([
    "重复控制不只比较文件名，应综合附件哈希、发票号码/代码、开票日期、金额、销售方和历史入账状态。",
    "保留原始电子凭证，不以截图或OCR文本替代原件；OCR数据属于辅助结构化信息。",
    "对凭证修改、转换、归档、移交和销毁实施权限及日志控制，并建立备份和恢复验证机制。",
    "对接电子凭证会计数据标准时，按凭证种类保存结构化数据、版式文件及元数据。",
])

doc.add_heading("11. 分期建设建议", level=1)
table(["阶段", "目标", "范围", "完成标志"], [
    ("一期：闭环可用", "完成线上报销主流程", "单据、唯一附件、OCR确认、提交、审批、退回、财务审核、基础报表、权限日志", "职工到财务审核全程线上；同名附件可独立管理"),
    ("二期：财务闭环", "打通付款与核算", "预算、付款批次、回单、ERP凭证接口、对账、规则配置", "审批通过至付款/入账可追踪并可对账"),
    ("三期：合规智能", "提高自动化和审计能力", "发票验真、电子凭证标准、档案接口、风险评分、智能审核、管理看板", "高频合规项自动检查，归档包可验证与移交"),
], [1400, 1900, 4100, 1960])

doc.add_heading("12. 一期验收标准", level=1)
bullets([
    "同一原文件名连续上传3次，应生成3个不同附件ID和唯一存储名；删除其中任意一条，其余两条仍存在且可预览。",
    "草稿可反复保存；提交后不能直接编辑，撤回或退回后按规则恢复编辑。",
    "审批路径按配置正确命中；所有同意、退回、拒绝和撤回动作均可追溯。",
    "系统能识别确定的重复发票并阻断或按配置升级处理，不以文件名作为唯一判断依据。",
    "财务可完成审核清单、科目/成本归属、付款状态和凭证号登记。",
    "普通职工不能查看他人无关单据和附件；附件预览、下载、删除均进行权限校验。",
    "关键操作日志完整，导出数据与页面统计口径一致；异常接口重试不产生重复单据。",
])

doc.add_heading("13. 建议运营指标", level=1)
table(["指标", "定义", "建议用途"], [
    ("平均提交时长", "从新建到首次提交的中位时长", "衡量职工填写体验"),
    ("一次通过率", "未退回即完成业务与财务审核的比例", "衡量规则前置与表单质量"),
    ("审批周期", "提交到财务审核通过的时长", "定位组织流程瓶颈"),
    ("财务单均处理时长", "财务有效操作时间/处理单量", "衡量自动化收益"),
    ("异常命中率", "重复、超标、缺件、预算等规则命中比例", "优化制度与规则"),
    ("付款及时率", "审批完成后在目标期限内付款比例", "衡量付款执行"),
    ("归档完整率", "具备完整凭证链和元数据的已入账单据比例", "衡量审计与档案质量"),
], [1700, 4800, 2860])

doc.add_heading("14. 待公司确认事项", level=1)
bullets([
    "费用类型、报销标准、发票抬头、可报期限及特殊场景材料清单。",
    "审批链及金额阈值、会签/加签/委托/撤回规则和财务主管复核条件。",
    "预算控制采用提示还是强控，以及部门、项目、个人等预算维度。",
    "付款方式、银企直联范围、制单复核分离要求和付款时效。",
    "ERP、预算、组织、企业IM、发票服务与档案系统的接口边界。",
    "电子档案保管范围、期限、移交格式、备份和销毁审批要求。",
])

doc.add_heading("15. 参考依据与调研来源", level=1)
refs = [
    "财政部、国家档案局：《关于规范电子会计凭证报销入账归档的通知》（财会〔2020〕6号），https://kjs.mof.gov.cn/zhengcefabu/202003/t20200331_3490938.htm",
    "财政部等部门：《关于推广应用电子凭证会计数据标准的通知》（财会〔2025〕9号），https://www.mof.gov.cn/jrttts/202505/t20250521_3964264.htm",
    "财政部会计司：《电子凭证会计数据标准（推广应用版）》，https://kjs.mof.gov.cn/zt/kuaijixinxihuajianshe/dzpzkjsjbzshsd/sjbz/202505/t20250519_3964020.htm",
    "国家档案局：《会计档案管理办法》修订内容要点，https://www.saac.gov.cn/daj/fzgz/201603/fb206b147418445d8cb319b711e9a51f.shtml",
    "国家档案局：《电子档案管理办法》发布信息，https://www.saac.gov.cn/daj/tzgg/202411/c5d21153b2cb47659a3ef5e0d1dc421c.shtml",
]
for r in refs:
    para = doc.add_paragraph(style="List Number")
    set_font(para.add_run(r), 9.2)
p("说明：本说明书吸收了公开产品常见的OCR、发票查重、预算控制、流程审批、付款、凭证和电子归档能力，但未将任何单一厂商的实现作为强制方案。")

doc.core_properties.title = "企业费用报销业务管理系统产品功能规划说明书"
doc.core_properties.subject = "职工报销、审批、财务审核、付款、入账与归档产品规划"
doc.core_properties.author = "产品规划组"
doc.save(OUT)
print(OUT.resolve())
