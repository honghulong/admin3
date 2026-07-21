from __future__ import annotations

from docx import Document
from pathlib import Path


def emit_paragraph(prefix: str, paragraph) -> None:
    text = paragraph.text.strip().replace("\t", "[TAB]")
    if text:
        print(f"{prefix}: {text}")


def walk_table(table, prefix: str) -> None:
    for row_index, row in enumerate(table.rows, start=1):
        for cell_index, cell in enumerate(row.cells, start=1):
            for paragraph_index, paragraph in enumerate(cell.paragraphs, start=1):
                emit_paragraph(f"{prefix} R{row_index}C{cell_index}P{paragraph_index}", paragraph)
            for nested_index, nested in enumerate(cell.tables, start=1):
                walk_table(nested, f"{prefix} R{row_index}C{cell_index}T{nested_index}")


document = Document("reference-original.docx")
lines: list[str] = []


def capture(prefix: str, paragraph) -> None:
    text = paragraph.text.strip().replace("\t", "[TAB]")
    if text:
        lines.append(f"{prefix}: {text}")


def capture_table(table, prefix: str) -> None:
    for row_index, row in enumerate(table.rows, start=1):
        for cell_index, cell in enumerate(row.cells, start=1):
            for paragraph_index, paragraph in enumerate(cell.paragraphs, start=1):
                capture(f"{prefix} R{row_index}C{cell_index}P{paragraph_index}", paragraph)
            for nested_index, nested in enumerate(cell.tables, start=1):
                capture_table(nested, f"{prefix} R{row_index}C{cell_index}T{nested_index}")


lines.extend((f"PARAGRAPHS={len(document.paragraphs)}", f"TABLES={len(document.tables)}"))
for index, paragraph in enumerate(document.paragraphs, start=1):
    capture(f"P{index}", paragraph)
for index, table in enumerate(document.tables, start=1):
    capture_table(table, f"T{index}")
for section_index, section in enumerate(document.sections, start=1):
    for label, part in (("HEADER", section.header), ("FOOTER", section.footer)):
        for paragraph_index, paragraph in enumerate(part.paragraphs, start=1):
            capture(f"S{section_index}-{label}-P{paragraph_index}", paragraph)
        for table_index, table in enumerate(part.tables, start=1):
            capture_table(table, f"S{section_index}-{label}-T{table_index}")

Path("source-inventory.txt").write_text("\n".join(lines), encoding="utf-8")
