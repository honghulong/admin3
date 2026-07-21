from __future__ import annotations

from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document


SOURCE = Path("reference-original.docx")
OUTPUT = Path("项目开发合同模板.docx")


def set_text(paragraph, value: str) -> None:
    """Replace text while retaining the paragraph's existing run formatting."""
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def replace_known_text(paragraph, replacements: dict[str, str]) -> None:
    text = paragraph.text
    for original, replacement in replacements.items():
        text = text.replace(original, replacement)
    if text != paragraph.text:
        set_text(paragraph, text)


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        yield from iter_table_paragraphs(table)
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
        for table in section.header.tables:
            yield from iter_table_paragraphs(table)
        for table in section.footer.tables:
            yield from iter_table_paragraphs(table)


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


document = Document(SOURCE)

# Fields that identify the original parties, project or commercial terms.
replacements = {
    "广东省东莞市人民检察院侵害未成年人案件强制报告管理系统建设（2026年）项目": "【项目名称】",
    "东莞市人民检察院侵害未成年人案件强制报告管理系统": "【项目名称】",
    "广东省东莞市人民检察院": "【甲方名称】",
    "广东维信智联科技有限公司": "【乙方名称】",
    "东莞市莞城街道兴塘社区东城西路185号": "【甲方地址】",
    "东莞市常平镇常平金塘路5号1栋204室": "【乙方地址】",
    "东莞市莞城区汇峰中心E区601": "【乙方通讯地址】",
    "沈丙友": "【甲方法定代表人】",
    "高旭东": "【甲方项目联系人】",
    "杨吉伟": "【乙方法定代表人】",
    "许柳飞": "【乙方项目联系人】",
    "13669813770": "【乙方联系电话】",
    "076922386992": "【乙方电话】",
    "076922386997": "【乙方传真】",
    "451207270@qq.com": "【乙方电子邮箱】",
    "中国银行东莞常平支行营业部": "【乙方开户银行】",
    "636673614878": "【乙方银行账号】",
    "东莞市": "【签订地点】",
}

for paragraph in iter_paragraphs(document):
    replace_known_text(paragraph, replacements)

# Replace project-specific body text but preserve clause positions and numbering.
body_slots = {
    2: "合同编号：【填写】",
    9: "项目名称：【项目名称】",
    10: "委托方（甲方）：【甲方名称】",
    11: "受托方（乙方）：【乙方名称】",
    12: "签订时间：【签订日期】",
    13: "签订地点：【签订地点】",
    14: "有效期限：【合同有效期限】",
    45: "委托方（甲方）：【甲方名称】",
    46: "住  所  地：【甲方地址】",
    47: "法定代表人：【甲方法定代表人】",
    48: "项目联系人：【甲方项目联系人】",
    49: "联系方式：【甲方联系电话】",
    50: "通讯地址：【甲方通讯地址】",
    51: "电    话：【甲方电话】    传    真：【甲方传真】",
    52: "电子信箱：【甲方电子邮箱】",
    53: "受托方（乙方）：【乙方名称】",
    54: "住  所  地：【乙方地址】",
    55: "法定代表人：【乙方法定代表人】",
    56: "项目联系人：【乙方项目联系人】",
    57: "联系方式：【乙方联系电话】",
    58: "通讯地址：【乙方通讯地址】",
    59: "电    话：【乙方电话】    传    真：【乙方传真】",
    60: "电子信箱：【乙方电子邮箱】",
    62: "本合同甲方委托乙方研究开发【项目名称】，并支付研究开发经费和报酬，乙方接受委托并进行研究开发工作。双方依据有关法律法规，经协商一致，订立本合同。",
    64: "1．技术目标：【填写本项目的技术目标、业务目标和预期成效】",
    65: "2．技术内容：【填写本项目的建设内容、功能范围、接口及交付物】",
    66: "3．技术方法和路线：【填写技术架构、实施路径、部署环境及安全要求】",
    67: "第二条  乙方应在本合同生效后【填写】日内向甲方提交研究开发计划。研究开发计划应包括以下主要内容：",
    68: "1．【填写阶段或工作项】；",
    69: "2．【填写阶段或工作项】；",
    70: "3．【填写阶段或工作项】；",
    71: "4．【填写阶段或工作项】。",
    73: "1．【填写阶段名称、计划工期和阶段成果】；",
    74: "2．【填写阶段名称、计划工期和阶段成果】；",
    75: "3．【填写阶段名称、计划工期和阶段成果】；",
    76: "4．【填写阶段名称、计划工期和阶段成果】；",
    77: "5．【填写阶段名称、计划工期和阶段成果】；",
    78: "6．【填写阶段名称、计划工期和阶段成果】；",
    79: "7．【填写其他项目进度要求】。",
    81: "1．技术资料清单：【填写甲方需提供的资料、数据、账号或其他条件】。",
    82: "2．提供时间和方式：【填写】。",
    83: "3．其他协作事项：【填写】。",
    84: "本合同履行完毕后，上述技术资料按以下方式处理：【填写】。",
    86: "1．研究开发经费和报酬总额为：人民币【合同总金额（大写）】（小写：￥【合同总金额】）。",
    87: "其中：【填写费用构成及比例】。",
    88: "2．研究开发经费由甲方【一次/分期/提成】支付乙方。具体支付方式和时间如下：",
    89: "（1）【填写付款节点、比例、金额及条件】；",
    90: "（2）【填写付款节点、比例、金额及条件】。",
    91: "付款前，乙方应向甲方提供【填写发票或结算材料要求】。",
    92: "乙方开户银行名称、地址和账号为：",
    93: "开户银行：【乙方开户银行】",
    94: "地址：【乙方开户行地址】",
    95: "账号：【乙方银行账号】",
    99: "请求，另一方应当在【填写】日内予以答复；逾期未予答复的，视为【填写】。",
    107: "的，双方按如下约定承担风险损失：【填写】。",
    112: "一方发现技术风险存在并可能导致研究开发失败或部分失败时，应当在【填写】日内通知另一方并采取适当措施减少损失。",
    114: "公开（包括以专利权方式公开），一方应在【填写】日内通知另一方解除合同。",
    118: "1．保密内容（包括技术信息和经营信息）：【填写】。",
    119: "2．涉密人员范围：【填写】。",
    120: "3．保密期限：【填写】。",
    121: "4．泄密责任：【填写】。",
    123: "1．保密内容（包括技术信息和经营信息）：【填写】。",
    124: "2．涉密人员范围：【填写】。",
    125: "3．保密期限：【填写】。",
    126: "4．泄密责任：【填写】。",
    128: "1．研究开发成果交付的形式及数量：【填写】。",
    129: "2．研究开发成果交付的时间及地点：时间【填写】；地点【填写】。",
    131: "果进行验收：【填写验收依据、标准、方法和验收期限】。",
    137: "专利权取得后的使用和有关利益分配方式如下：【填写知识产权归属及利益分配约定】。",
    140: "（1）技术秘密的使用权：【填写】；（2）技术秘密的转让权：【填写】；",
    141: "（3）相关利益的分配办法：【填写】。",
    142: "双方对本合同有关的知识产权权利归属特别约定如下：【填写】。",
    150: "1．技术服务和指导内容：【填写】。",
    151: "2．地点和方式：【填写】。",
    152: "3．售后服务：【填写服务期限、响应要求、维护范围及费用】。",
    155: "1．乙方违反本合同约定的违约责任：【填写】。",
    156: "2．甲方违反本合同约定的违约责任：【填写】。",
    157: "3．甲方违反保密义务的违约责任：【填写】。",
    158: "4．乙方违反保密义务的违约责任：【填写】。",
    159: "5．乙方违反成果转让限制的违约责任：【填写】。",
    160: "6．乙方违反技术服务义务的违约责任：【填写】。",
    163: "步特征的新的技术成果及其权利归属，由【甲方/乙方/双方】享有。",
    164: "具体相关利益的分配办法如下：【填写】。",
    167: "的技术成果，归【甲方/乙方/双方】所有。具体相关利益的分配办法如下：【填写】。",
    168: "【填写】。",
    169: "第二十二条  双方确定，在本合同有效期内，甲方指定【甲方项目联系人】为甲方项目联系人，乙方指定【乙方项目联系人】为乙方项目联系人。",
    180: "1．提交【填写仲裁委员会】仲裁；",
    181: "2．依法向【填写有管辖权的人民法院】起诉。",
    191: "1．技术背景资料：【填写】；2．可行性论证报告：【填写】；",
    192: "3．技术评价报告：【填写】；4．技术标准和规范：【填写】；",
    193: "5．原始设计和工艺文件：【填写】；6．其他：【填写】。",
    194: "第二十七条  双方约定本合同其他相关事项为：【填写】。",
    195: "第二十八条  本合同一式【填写】份，甲方执【填写】份，乙方执【填写】份，具有同等法律效力。",
    197: "附件1：【项目功能/服务清单】",
    198: "附件2：保密协议",
    202: "甲方：【甲方名称】                    （盖章）",
    203: "法定代表人/委托代理人：                           （签名）",
    204: "签署日期：      年     月     日",
    206: "乙方：【乙方名称】                    （盖章）",
    207: "法定代表人/委托代理人：                           （签名）",
    208: "签署日期：      年     月     日",
    210: "附件1：【项目功能/服务清单】",
    219: "甲方：【甲方名称】",
    220: "乙方：【乙方名称】",
    223: "本保密协议中所涉及的保密信息指【项目名称】所涉及的信息资料，包括但不限于：用户需求、数据、信息、协议、参考资料、内部发文、通知、会议记录、往来函件和电子邮件等。",
    248: "对因乙方主观故意或重大过失引起的泄密事件承担相应法律责任。违约责任、赔偿范围及处理方式由双方在本协议或补充协议中约定：【填写】。",
    251: "3.2   本保密协议为双方签订的《【项目名称】》的合同附件；双方合同终止不影响本协议约定的保密义务。",
}

for paragraph_index, value in body_slots.items():
    set_text(document.paragraphs[paragraph_index - 1], value)

# Keep the annex table layout while removing the original project's functional detail.
for table in document.tables:
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            continue
        for cell_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                if cell_index == 0:
                    set_text(paragraph, "")
                else:
                    set_text(paragraph, "【填写】")

properties = document.core_properties
properties.author = ""
properties.last_modified_by = ""
properties.title = "技术开发（委托）合同模板"
properties.subject = ""
properties.keywords = ""
properties.comments = ""

document.save(OUTPUT)

# Text boxes and similar DrawingML content are not exposed by python-docx.
# Apply the same targeted scrubbing to every XML part in the finished package.
with ZipFile(OUTPUT, "r") as source, NamedTemporaryFile(
    delete=False, suffix=".docx", dir=OUTPUT.parent
) as temp_file:
    temp_path = Path(temp_file.name)
    with ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                # Source text boxes can overlap the form fields above. They must
                # be blanked rather than converted into a second visible slot.
                for original in ("沈丙友", "高旭东", "杨吉伟", "许柳飞"):
                    content = content.replace(original.encode("utf-8"), b"")
                for original, replacement in replacements.items():
                    content = content.replace(original.encode("utf-8"), replacement.encode("utf-8"))
            target.writestr(item, content)
temp_path.replace(OUTPUT)

print(f"Created {OUTPUT}")
