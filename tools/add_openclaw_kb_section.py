from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"D:\hongmengProjects\admin3\AI开发工具与本地智能体应用总结2026.docx")
BACKUP_PATH = DOCX_PATH.with_suffix(".before_openclaw_kb_section.docx")


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def insert_table_after(paragraph: Paragraph, rows: list[list[str]]):
    parent = paragraph._parent
    table = parent.add_table(rows=len(rows), cols=len(rows[0]), width=Cm(15.2))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph._p.addnext(table._tbl)
    return table


def insert_paragraph_after_table(table, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    p = Paragraph(new_p, table._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document(DOCX_PATH)

    # Remove an existing 1.4 section if this script is rerun.
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("1.4"):
            start = i
        elif start is not None and p.style.name.startswith("Heading") and not p.text.strip().startswith("1.4"):
            end = i
            break
    if start is not None:
        if end is None:
            end = len(doc.paragraphs)
        for p in doc.paragraphs[start:end]:
            p._element.getparent().remove(p._element)

    # Re-load after deletion to keep paragraph references fresh.
    doc.save(DOCX_PATH)
    doc = Document(DOCX_PATH)

    insert_after = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("1.3"):
            insert_after = p
        elif insert_after is not None and p.style.name == "Heading 2":
            break
        elif insert_after is not None:
            insert_after = p

    if insert_after is None:
        raise RuntimeError("未找到 1.3 OCR 服务章节")

    p = insert_paragraph_after(insert_after, "1.4 OpenClaw 本地知识库", "Heading 3")
    p = insert_paragraph_after(
        p,
        "定位：OpenClaw 的本地知识库是面向本机业务资料的 RAG 检索能力，解决“资料很多但人工翻找慢”的问题。"
        "它不是把全部文件直接塞进对话上下文，而是先把业务文档解析、切片、向量化并入库；用户在 OpenClaw、ClawX 或桥接服务中提问时，"
        "系统先召回相关片段，再把片段作为上下文交给大模型生成回答。",
    )
    p = insert_paragraph_after(
        p,
        "数据范围：当前配置以本地知识库工具目录的 sources.txt 和 config_v2.py 为准，覆盖 D:\\项目信息、F:\\B部门资料、F:\\C软件产品、"
        "F:\\x项目文档、F:\\w维信智联、F:\\c财政预算 等业务资料目录。系统只索引文档型资料，支持 txt、md、doc/docx、pdf、pptx、"
        "xlsx/xls、html、xml、json、yaml、csv 等格式，并排除 .git、node_modules、虚拟环境、缓存文件和 Office 临时文件。",
    )
    p = insert_paragraph_after(
        p,
        "索引流程：定时任务先用 Everything SDK 的 es.exe 秒级扫描候选文件；如果在定时任务或服务会话里 es.exe IPC 不可用，脚本会回退到 os.walk。"
        "每个文件用“文件大小+修改时间”生成轻量指纹，只有新增、修改或删除的文件才进入增量处理。解析阶段按格式提取文本：PDF 优先读取文本层，"
        "扫描件可按配置调用本地 OCR 服务；DOCX 通过 zip/XML 读取正文，老格式或异常文件可走兼容解析；PPTX、XLSX 等也各自提取可读文本。",
    )
    p = insert_paragraph_after(
        p,
        "向量化与存储：文本使用 RecursiveCharacterTextSplitter 按 500 字符、100 重叠切片，再用本地 sentence-transformers 模型 "
        "BAAI/bge-small-zh-v1.5 在 CPU 上生成 512 维 embedding。早期版本使用 Chroma 向量库，后来因为 cron 超时或进程被杀时容易造成 "
        "sqlite WAL 与 HNSW 索引损坏，已迁移为 MySQL 8.4 方案：documents 表保存文件元数据和索引状态，chunks 表保存切片正文与 float32 "
        "向量 BLOB，file_blacklist 记录空文件、解析失败和超时文件，process_logs/run_stats/daily_health 用于审计和健康统计。",
    )
    p = insert_paragraph_after(
        p,
        "查询流程：语义检索脚本 query_kb_mysql.py 会把问题编码成同一模型的向量，从 MySQL 读取 chunks 的 embedding，转为 numpy float32 数组后批量计算余弦相似度，返回 Top K 片段及来源文件。"
        "日常快速查文件或关键词时，还提供 kb_search.py，不加载 embedding 模型，直接在 MySQL chunks/documents 上做关键词检索，适合先定位资料，再做语义问答。",
    )

    table_rows = [
        ["组成", "实现细节", "作用"],
        ["数据源", "sources.txt / config_v2.py 管理 D 盘和 F 盘业务目录", "限定知识库覆盖范围，避免把代码和缓存混入业务资料"],
        ["扫描", "Everything es.exe；失败时回退 os.walk", "快速发现新增、修改、删除文件"],
        ["解析", "按文件类型读取文本；扫描型 PDF 可选调用本地 OCR", "把文档、表格、演示稿和 PDF 转成可切片文本"],
        ["切片与向量", "500 字符切片、100 重叠；BAAI/bge-small-zh-v1.5；512 维向量", "形成可检索的语义片段"],
        ["存储", "MySQL 8.4：documents、chunks、file_blacklist、process_logs、run_stats、daily_health", "提高定时任务稳定性和可运维性"],
        ["查询入口", "query_kb_mysql.py、kb_search.py、Hermes Bridge /kb/query", "供 OpenClaw、ClawX 或外部智能体调用本地知识库"],
        ["定时任务", "ClawX/OpenClaw cron 每天 13:30 触发增量更新；Hermes-Win 记录显示脚本模式已运行多次", "让资料库随本地文件变化自动刷新"],
    ]
    table = insert_table_after(p, table_rows)

    p = insert_paragraph_after_table(
        table,
        "运行与调度：ClawX/OpenClaw 的 cron 配置中有“知识库增量更新”任务，Asia/Shanghai 时区每天 13:30 执行，早期命令为进入 "
        "OpenClaw 工作区的 knowledge-base 目录后运行 build_index_mysql.py。Hermes 迁移后也保留同名任务，脚本模式直接调用知识库更新脚本，"
        "最近任务状态显示为 ok，说明该链路已进入自动维护状态。",
    )
    insert_paragraph_after(
        p,
        "经验结论：这套实现的关键价值不是单个模型，而是把“文件发现、文本提取、OCR 补充、增量索引、向量召回、运行统计”串成了可复用基础能力。"
        "后续接入新的办公资料、合同、项目文档或制度库时，只需要把源目录纳入配置并保持 13:30 增量任务运行，即可让 OpenClaw 在对话中调用最新本地知识。",
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
