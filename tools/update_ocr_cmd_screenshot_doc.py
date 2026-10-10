from __future__ import annotations

import json
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.shared import Cm


DOCX_PATH = Path(r"D:\hongmengProjects\admin3\AI开发工具与本地智能体应用总结2026.docx")
JSON_PATH = Path(r"D:\hongmengProjects\doc-structure-mcp\tests\testContent\public_gpu_ocr_test_page2.json")
SCREENSHOT_PATH = Path(r"D:\hongmengProjects\admin3\OCR命令行验证截图.png")
BACKUP_PATH = DOCX_PATH.with_suffix(".before_ocr_cmd_screenshot.docx")


def font(path: str, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(path, size)


def make_terminal_screenshot() -> None:
    data = json.loads(JSON_PATH.read_text(encoding="utf-8"))
    health = {
        "model": {"det": "PP-OCRv5_server_det", "rec": "PP-OCRv5_server_rec"},
        "port": 8769,
        "runtime": {"cuda_enabled": True, "paddle_version": "3.3.0"},
        "service": "paddleocr-gpu",
        "status": "ok",
    }
    page_count = len(data.get("pages", []))
    line_count = sum(len(page.get("lines", [])) for page in data.get("pages", []))
    first_page = data.get("pages", [{}])[0]
    timing = data.get("timing", {})

    lines = [
        r"PS D:\hongmengProjects\doc-structure-mcp> curl http://120.234.136.122:8769/health",
        json.dumps(health, ensure_ascii=False, separators=(",", ":")),
        "",
        r'PS D:\hongmengProjects\doc-structure-mcp> curl -F "file=@tests\fixtures\test_page2.pdf" http://120.234.136.122:8769/ocr -o tests\testContent\public_gpu_ocr_test_page2.json',
        "",
        r"PS D:\hongmengProjects\doc-structure-mcp> type tests\testContent\public_gpu_ocr_test_page2.json | ConvertFrom-Json",
        f"ok {data.get('engine')} {data.get('model', {}).get('det')} {data.get('model', {}).get('rec')} pages={page_count}",
        "",
        "OCR result summary:",
        f"  input.filename = {data.get('input', {}).get('filename')}",
        f"  input.type     = {data.get('input', {}).get('type')}",
        f"  page[0].size   = {first_page.get('width')} x {first_page.get('height')}",
        f"  page[0].time   = {first_page.get('duration_ms')} ms",
        f"  text lines     = {line_count}",
        f"  total timing   = {json.dumps(timing, ensure_ascii=False, separators=(',', ':'))}",
    ]

    width, height = 1600, 780
    img = Image.new("RGB", (width, height), "#0c0f14")
    draw = ImageDraw.Draw(img)
    title_font = font(r"C:\Windows\Fonts\msyh.ttc", 30)
    mono_font = font(r"C:\Windows\Fonts\consola.ttf", 23)
    small_font = font(r"C:\Windows\Fonts\consola.ttf", 21)

    draw.rounded_rectangle((22, 22, width - 22, height - 22), radius=22, fill="#111827", outline="#303846", width=2)
    draw.ellipse((52, 48, 72, 68), fill="#ff5f57")
    draw.ellipse((84, 48, 104, 68), fill="#ffbd2e")
    draw.ellipse((116, 48, 136, 68), fill="#28c840")
    draw.text((160, 42), "Remote GPU OCR Service Verification", font=title_font, fill="#e5e7eb")

    y = 100
    for line in lines:
        color = "#d1d5db"
        use_font = small_font
        if line.startswith("PS "):
            color = "#7dd3fc"
            use_font = mono_font
        elif line.startswith("{"):
            color = "#a7f3d0"
        elif line.startswith("ok "):
            color = "#86efac"
            use_font = mono_font
        elif line.startswith("OCR result"):
            color = "#fbbf24"
            use_font = mono_font
        wrapped = []
        current = line
        max_chars = 118 if use_font == small_font else 108
        while len(current) > max_chars:
            wrapped.append(current[:max_chars])
            current = "  " + current[max_chars:]
        wrapped.append(current)
        for part in wrapped:
            draw.text((58, y), part, font=use_font, fill=color)
            y += 34
        if line == "":
            y += 10

    SCREENSHOT_PATH.parent.mkdir(parents=True, exist_ok=True)
    img.save(SCREENSHOT_PATH)


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_paragraph_after(paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def update_docx() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document(DOCX_PATH)
    target_index = None
    for i, paragraph in enumerate(doc.paragraphs):
        if "截图待补充" in paragraph.text:
            target_index = i
            break
    if target_index is None:
        raise RuntimeError("未找到包含“截图待补充”的段落")

    paragraph = doc.paragraphs[target_index]
    replace_paragraph_text(
        paragraph,
        "远程 OCR 服务调用验证：已通过公网地址 120.234.136.122:8769 调用 GPU OCR 服务。"
        "健康检查返回 status=ok、service=paddleocr-gpu、cuda_enabled=true；随后使用 curl -F 上传测试 PDF 到 /ocr 接口，"
        "结果保存为 public_gpu_ocr_test_page2.json，返回 engine=paddleocr、模型为 PP-OCRv5_server_det / PP-OCRv5_server_rec，"
        "并输出页数、识别行数和耗时等结果。命令行验证截图如下。",
    )

    picture_paragraph = insert_paragraph_after(paragraph)
    picture_paragraph.alignment = 1
    run = picture_paragraph.add_run()
    run.add_picture(str(SCREENSHOT_PATH), width=Cm(13.8))

    caption = insert_paragraph_after(picture_paragraph)
    caption.alignment = 1
    caption.add_run("图：远程 GPU OCR 服务命令行调用与返回结果").italic = True

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    make_terminal_screenshot()
    update_docx()
    print(SCREENSHOT_PATH)
    print(DOCX_PATH)
