from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"D:\hongmengProjects\admin3\AI开发工具与本地智能体应用总结2026.docx")
BACKUP_PATH = DOCX_PATH.with_suffix(".before_honghuang_voice_section.docx")


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def insert_table_after(paragraph: Paragraph, rows: list[list[str]]):
    parent = paragraph._parent
    table = parent.add_table(rows=len(rows), cols=len(rows[0]), width=Cm(15.2))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = text
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph._p.addnext(table._tbl)
    return table


def insert_paragraph_after_table(table, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    p = Paragraph(new_p, table._parent)
    if style:
        p.style = style
    if text:
        p.add_run(text)
    return p


def remove_existing_section(doc: Document, section_prefix: str) -> None:
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text.startswith(section_prefix):
            start = i
        elif start is not None and p.style.name.startswith("Heading") and not text.startswith(section_prefix):
            end = i
            break

    if start is None:
        return
    if end is None:
        end = len(doc.paragraphs)
    for p in doc.paragraphs[start:end]:
        p._element.getparent().remove(p._element)


def find_insert_after(doc: Document) -> Paragraph:
    insert_after = None
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith("1.4"):
            insert_after = p
        elif insert_after is not None and p.style.name == "Heading 2":
            break
        elif insert_after is not None:
            insert_after = p

    if insert_after is None:
        raise RuntimeError("未找到 1.4 OpenClaw 本地知识库章节")
    return insert_after


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    doc = Document(DOCX_PATH)
    remove_existing_section(doc, "1.5")
    doc.save(DOCX_PATH)

    doc = Document(DOCX_PATH)
    insert_after = find_insert_after(doc)

    p = insert_paragraph_after(insert_after, "1.5 鸿蒙语音助手原型", "Heading 3")
    p = insert_paragraph_after(
        p,
        "定位：HonghuangVoice 是使用 TRAE 开发的 HarmonyOS 语音助手原型，主要用于验证鸿蒙端能否把语音交互、"
        "本地智能体和 OpenClaw 调用串起来。它更接近技术预研和演示样机，不是已经形成稳定需求、稳定用户和生产流程的业务系统。",
    )
    p = insert_paragraph_after(
        p,
        "项目形态：代码目录包含 AppScope、entry、hvigor、oh-package.json5、build-profile.json5 等 HarmonyOS 工程结构，"
        "入口页面和服务层采用 ArkTS 编写。README 和操作指引中记录了语音识别、语音合成、AI 对话、全双工对话、唤醒模式等功能；"
        "代码中也能看到 Index.ets、MainViewModel.ets、ManualRecordService.ets、FullDuplexConversationService.ets、"
        "HonghuangWakeupService.ets、HonghuangAIService.ets 等模块分工。",
    )
    p = insert_paragraph_after(
        p,
        "已验证思路：前端提供手动录音、文本输入、全双工对话、唤醒模式四类入口。语音侧使用 CoreSpeechKit 相关能力做 ASR/TTS，"
        "并通过 VAD、状态机、TTS 打断、请求编号和过期回复丢弃等方式，模拟“持续听、可打断、再回答”的对话体验。AI 侧通过 "
        "OpenAI 兼容的 /v1/chat/completions HTTP 接口访问本地或远程 OpenClaw 服务，模型配置为 openclaw/default。",
    )

    table_rows = [
        ["模块", "已经做过的验证", "现实边界"],
        ["鸿蒙端工程", "ArkTS 页面、ViewModel、服务层和资源文件已形成基本工程结构", "仍需真机、签名和系统权限配合，不能只靠模拟器判断体验"],
        ["语音输入输出", "手动录音、ASR 识别、TTS 播报、文本输入复用播报链路", "识别质量、延迟和播报稳定性依赖设备与系统语音能力"],
        ["全双工/打断", "通过 VAD、状态机、TTS stop、请求编号实现应用侧近似全双工", "华为系统级小艺真人感全双工能力未向普通第三方 App 直接开放"],
        ["唤醒模式", "设置“洪荒洪荒”等唤醒词变体，唤醒后进入指令识别和播报流程", "长期后台监听、功耗、误唤醒和权限策略需要专门评估"],
        ["OpenClaw 接入", "通过 OpenAI 兼容 HTTP 接口调用 OpenClaw，对接本地智能体能力", "依赖 OpenClaw 服务在线、网络可达和本地配置，响应速度也可能偏慢"],
        ["落地条件", "可作为本地智能体语音入口的探索样机保留", "当前缺少明确业务场景、目标用户和使用频率，因此不建议包装成已落地应用"],
    ]
    table = insert_table_after(p, table_rows)

    p = insert_paragraph_after_table(
        table,
        "局限性：项目资料中已经判断，小艺 Claw 的底层能力和系统级全双工语音能力并没有以通用 API 形式直接开放给普通第三方 App。"
        "因此 HonghuangVoice 的路线是“鸿蒙 App 自己做语音交互 + 通过 HTTP 调 OpenClaw”，能验证思路，但还不能等同于系统级语音助手能力。",
    )
    insert_paragraph_after(
        p,
        "管理结论：该原型可以作为 TRAE 开发鸿蒙应用、语音交互和本地智能体接入的技术储备保留。后续如果要继续推进，应先明确真实需求，"
        "例如内部资料语音问答、移动端任务录入、现场助手或特定业务流程入口；在没有明确需求前，建议不投入上线包装和长期运维，只保留代码、"
        "操作说明和关键技术结论。",
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
