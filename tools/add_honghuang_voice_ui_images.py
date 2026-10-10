from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(r"D:\hongmengProjects\admin3\AI开发工具与本地智能体应用总结2026.docx")
BACKUP_PATH = DOCX_PATH.with_suffix(".before_honghuang_voice_ui_images.docx")
OUT_DIR = Path(r"D:\hongmengProjects\admin3\deploy_tmp\honghuang_voice_ui")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


F_TITLE = font(32, True)
F_H2 = font(24, True)
F_TEXT = font(20)
F_SMALL = font(16)
F_TINY = font(14)


def rounded(draw: ImageDraw.ImageDraw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text(draw: ImageDraw.ImageDraw, xy, value, fill="#ffffff", fnt=F_TEXT, anchor=None):
    draw.text(xy, value, fill=fill, font=fnt, anchor=anchor)


def gradient_bg(w: int, h: int) -> Image.Image:
    img = Image.new("RGB", (w, h), "#1a1a2e")
    pix = img.load()
    top = (26, 26, 46)
    mid = (22, 33, 62)
    bot = (15, 52, 96)
    for y in range(h):
        t = y / max(1, h - 1)
        if t < 0.5:
            tt = t / 0.5
            c = tuple(int(top[i] * (1 - tt) + mid[i] * tt) for i in range(3))
        else:
            tt = (t - 0.5) / 0.5
            c = tuple(int(mid[i] * (1 - tt) + bot[i] * tt) for i in range(3))
        for x in range(w):
            pix[x, y] = c
    return img


def draw_header(draw: ImageDraw.ImageDraw, selected: str, wake_on: bool = False):
    text(draw, (44, 62), "洪荒", fnt=F_TITLE)
    rounded(draw, (238, 52, 392, 88), 14, "#ff6b6b")
    text(draw, (315, 70), "停止播报", fnt=F_SMALL, anchor="mm")
    rounded(draw, (410, 52, 590, 88), 12, "#2d3a4a")
    text(draw, (500, 70), selected, fnt=F_TINY, anchor="mm")
    rounded(draw, (610, 50, 674, 80), 15, "#ff6b35" if wake_on else "#4a5568")
    knob_x = 650 if wake_on else 634
    draw.ellipse((knob_x - 12, 53, knob_x + 12, 77), fill="#ffffff")
    text(draw, (642, 102), "唤醒模式", "#ff6b35" if wake_on else "#aaaaaa", F_TINY, "mm")


def draw_message(draw: ImageDraw.ImageDraw, role: str, body: str, y: int):
    if role == "user":
        x1, x2 = 250, 672
        rounded(draw, (x1, y, x2, y + 78), 18, "#95ec69")
        text(draw, (x1 + 24, y + 22), body, "#ffffff", F_TEXT)
    else:
        x1, x2 = 40, 514
        rounded(draw, (x1, y, x2, y + 96), 18, "#ffffff")
        for i, line in enumerate(body.split("\n")):
            text(draw, (x1 + 24, y + 18 + i * 27), line, "#111111", F_SMALL)
        draw.ellipse((528, y + 34, 572, y + 78), fill="#eef2f7")
        text(draw, (550, y + 55), "▶", "#777777", F_SMALL, "mm")


def draw_footer(draw: ImageDraw.ImageDraw, input_text: str = "", focus: bool = True, full_duplex: bool = False):
    rounded(draw, (34, 914, 578, 968), 27, "#2d3a4a")
    text(draw, (60, 941), input_text or "输入消息...", "#ffffff" if input_text else "#888888", F_SMALL, "lm")
    rounded(draw, (594, 914, 680, 968), 27, "#4ecdc4")
    text(draw, (637, 941), "发送", "#ffffff", F_SMALL, "mm")
    if focus:
        if full_duplex:
            rounded(draw, (246, 990, 506, 1058), 34, "#ff6b6b")
            text(draw, (376, 1024), "停止全双工", "#ffffff", F_TEXT, "mm")
            text(draw, (376, 1086), "全双工模式：持续对话，随时打断", "#7c4dff", F_TINY, "mm")
        else:
            rounded(draw, (184, 990, 334, 1058), 34, "#4ecdc4")
            text(draw, (259, 1024), "录音", "#ffffff", F_TEXT, "mm")
            rounded(draw, (354, 990, 542, 1058), 34, "#7c4dff")
            text(draw, (448, 1024), "全双工对话", "#ffffff", F_TEXT, "mm")
            text(draw, (363, 1086), "点击按钮开始录音，再次点击停止", "#888888", F_TINY, "mm")


def draw_status(draw: ImageDraw.ImageDraw, label: str, dot: str, color: str, y: int = 820):
    text(draw, (360, y), label, "#cccccc" if color != "#4ecdc4" else "#4ecdc4", F_SMALL, "mm")
    draw.ellipse((304, y + 28, 316, y + 40), fill=color)
    text(draw, (362, y + 34), dot, "#888888", F_TINY, "mm")


def make_screen(path: Path, mode: str):
    img = gradient_bg(720, 1120)
    draw = ImageDraw.Draw(img)
    if mode == "idle":
        draw_header(draw, "智能体")
        text(draw, (360, 354), "暂无对话记录，点击下方按钮开始对话", "#777777", F_SMALL, "mm")
        draw_status(draw, "空闲", "空闲", "#888888")
        draw_footer(draw)
    elif mode == "chat":
        draw_header(draw, "openclaw+openai")
        draw_message(draw, "user", "帮我总结今天的会议纪要", 182)
        draw_message(draw, "assistant", "已整理三点：\n1. 预算审批待确认\n2. 周五前补齐材料", 292)
        draw_message(draw, "user", "再生成一版短摘要", 428)
        draw_status(draw, "回复中...", "回复中", "#ff6b6b")
        draw_footer(draw, "再生成一版短摘要")
    elif mode == "record":
        draw_header(draw, "openclaw+openai")
        draw_message(draw, "assistant", "我在听，请开始说话。", 220)
        for i, h in enumerate([36, 58, 84, 54, 72]):
            x = 300 + i * 28
            rounded(draw, (x, 598 - h // 2, x + 8, 598 + h // 2), 4, "#4ecdc4")
        draw_status(draw, "正在聆听...", "聆听中", "#4ecdc4")
        draw_footer(draw)
    elif mode == "full":
        draw_header(draw, "openclaw+openai")
        draw_message(draw, "user", "洪荒，继续说重点", 170)
        draw_message(draw, "assistant", "已暂停上一段播报，\n现在只保留结论。", 282)
        for i, h in enumerate([52, 86, 44, 76, 58]):
            x = 300 + i * 28
            rounded(draw, (x, 590 - h // 2, x + 8, 590 + h // 2), 4, "#7c4dff")
        draw_status(draw, "全双工对话中...", "全双工", "#7c4dff")
        draw_footer(draw, full_duplex=True)
    elif mode == "wake":
        draw_header(draw, "openclaw+openai", wake_on=True)
        text(draw, (360, 214), "等待唤醒词", "#ff6b35", F_H2, "mm")
        rounded(draw, (188, 270, 532, 402), 28, "#213044", "#ff6b35", 2)
        text(draw, (360, 318), "“洪荒洪荒”", "#ffffff", F_TITLE, "mm")
        text(draw, (360, 366), "唤醒后进入指令识别和播报流程", "#aaaaaa", F_SMALL, "mm")
        draw_status(draw, "等待唤醒词...", "唤醒", "#4ecdc4")
        draw_footer(draw)
    elif mode == "xiaoyi":
        rounded(draw, (26, 52, 96, 88), 10, "#2d3a4a")
        text(draw, (61, 70), "返回", "#ffffff", F_SMALL, "mm")
        text(draw, (176, 70), "小红虾智能体", "#ffffff", F_H2, "lm")
        rounded(draw, (54, 148, 666, 1016), 22, "#f8fafc")
        text(draw, (360, 204), "小红虾智能体", "#111827", F_H2, "mm")
        rounded(draw, (98, 264, 622, 356), 18, "#eef2ff")
        text(draw, (126, 294), "你好，我是小红虾智能体。", "#1f2937", F_SMALL)
        text(draw, (126, 322), "可通过系统智能体组件承载对话。", "#6b7280", F_TINY)
        rounded(draw, (98, 876, 536, 936), 30, "#e5e7eb")
        text(draw, (126, 906), "输入要咨询的问题...", "#6b7280", F_SMALL, "lm")
        rounded(draw, (554, 876, 622, 936), 30, "#4ecdc4")
        text(draw, (588, 906), "发送", "#ffffff", F_SMALL, "mm")
    elif mode == "office":
        rounded(draw, (26, 52, 96, 88), 10, "#2d3a4a")
        text(draw, (61, 70), "返回", "#ffffff", F_SMALL, "mm")
        text(draw, (176, 70), "日常办公助手", "#ffffff", F_H2, "lm")
        rounded(draw, (54, 148, 666, 1016), 22, "#f8fafc")
        text(draw, (360, 204), "日常办公助手", "#111827", F_H2, "mm")
        cards = [
            ("会议纪要", "整理录音、提炼待办事项"),
            ("材料起草", "按要求生成初稿和摘要"),
            ("任务跟踪", "记录待办、提醒下一步"),
        ]
        for idx, (title, desc) in enumerate(cards):
            y = 276 + idx * 132
            rounded(draw, (98, y, 622, y + 92), 18, "#ffffff", "#d1d5db", 2)
            text(draw, (126, y + 30), title, "#111827", F_TEXT)
            text(draw, (126, y + 62), desc, "#6b7280", F_SMALL)
        rounded(draw, (98, 876, 536, 936), 30, "#e5e7eb")
        text(draw, (126, 906), "输入办公任务...", "#6b7280", F_SMALL, "lm")
        rounded(draw, (554, 876, 622, 936), 30, "#4ecdc4")
        text(draw, (588, 906), "发送", "#ffffff", F_SMALL, "mm")
    img.save(path)


def insert_paragraph_after(paragraph: Paragraph, text_value: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style:
        p.style = style
    if text_value:
        p.add_run(text_value)
    return p


def find_section_end(doc: Document) -> Paragraph:
    target = None
    for p in doc.paragraphs:
        stripped = p.text.strip()
        if stripped.startswith("1.5"):
            target = p
        elif target is not None and p.style.name.startswith("Heading") and not stripped.startswith("1.5"):
            break
        elif target is not None:
            target = p
    if target is None:
        raise RuntimeError("未找到 1.5 鸿蒙语音助手原型章节")
    return target


def remove_existing_ui_block(doc: Document) -> None:
    start = end = None
    for i, p in enumerate(doc.paragraphs):
        stripped = p.text.strip()
        if stripped == "界面示意：":
            start = i
        elif start is not None and p.style.name.startswith("Heading"):
            end = i
            break
    if start is None:
        return
    if end is None:
        end = len(doc.paragraphs)
    for p in doc.paragraphs[start:end]:
        p._element.getparent().remove(p._element)


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    screens = [
        ("idle", "图 1.5-1 主界面空闲态：顶部智能体选择、唤醒开关、底部输入与录音入口。"),
        ("chat", "图 1.5-2 文本对话态：用户消息、AI 回复、单条消息播报入口。"),
        ("record", "图 1.5-3 手动录音态：ASR 聆听状态和语音波形反馈。"),
        ("full", "图 1.5-4 全双工态：持续监听、播报打断和停止全双工入口。"),
        ("wake", "图 1.5-5 唤醒模式态：唤醒词触发后进入指令识别。"),
        ("xiaoyi", "图 1.5-6 小红虾智能体页：通过系统智能体组件承载特定智能体对话。"),
        ("office", "图 1.5-7 日常办公助手页：面向办公任务的智能体入口示意。"),
    ]
    image_paths = []
    for mode, _caption in screens:
        path = OUT_DIR / f"honghuang_voice_{mode}.png"
        make_screen(path, mode)
        image_paths.append(path)

    doc = Document(DOCX_PATH)
    remove_existing_ui_block(doc)
    doc.save(DOCX_PATH)

    doc = Document(DOCX_PATH)
    p = find_section_end(doc)
    p = insert_paragraph_after(
        p,
        "界面示意：",
    )
    p = insert_paragraph_after(
        p,
        "以下图片为根据 HonghuangVoice 原型代码和操作指引绘制的界面示意，用于说明功能入口和交互状态；未作为真机运行截图使用。",
    )
    for path, (_mode, caption) in zip(image_paths, screens):
        p = insert_paragraph_after(p, "")
        run = p.add_run()
        run.add_picture(str(path), width=Cm(6.0))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = insert_paragraph_after(p, caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    for path in image_paths:
        print(path)


if __name__ == "__main__":
    main()
